"""python3 tools/review_sheets.py  ->  release/review/square-packing-review-<lang>.xlsx

One workbook per draft language, from release/review/strings.json (made by
tools/review_sheets.js).  Sheet "How to review": instructions, the key terms to check
first, and an example row.  Sheet "Strings": every string with its context, the English,
the draft, and two yellow columns for the reviewer.  No formulas.
"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent.parent
data = json.loads((ROOT / 'release' / 'review' / 'strings.json').read_text())

ARIAL = 'Arial'
HEAD = Font(name=ARIAL, bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='1F5F8B')
INPUT_FILL = PatternFill('solid', fgColor='FFF2A8')
MUTED = Font(name=ARIAL, color='5D6B66', size=9)
BODY = Font(name=ARIAL, size=10)
BOLD = Font(name=ARIAL, size=10, bold=True)
TITLE = Font(name=ARIAL, size=14, bold=True)
WRAP = Alignment(wrap_text=True, vertical='top')
THIN = Border(bottom=Side(style='thin', color='D3DBD8'))

TERMS = [  # (term, key whose draft shows the word)
    ('the game\'s name', 'app.name'), ('level', 'levels.title'), ('Free play', 'mode.free'),
    ('frame', 'goal.fit'), ('hint', 'btn.hint'), ('show me how (the answer)', 'btn.showAnswer'),
    ('start over', 'btn.reset'), ('skip', 'btn.skip'), ('tray', 'help.double'),
    ('packing', 'move.eyebrow'), ('standard packing', 'btn.compare'), ('puzzle', 'free.newRandom'),
    ('Advanced', 'adv.title'), ('strategy', 'ans.ledger'),
    ('floor (L* without the smallest square)', 'ans.floorRow'), ('strategy names', None),
]
INSTRUCTIONS = [
    'Thank you for reviewing this translation of Square Packing, a browser puzzle for a general audience about fitting squares into the smallest square frame. Plain, friendly wording matters more than technical precision, except in the Advanced answer panel.',
    'Each row of the "Strings" sheet is one piece of text from the game: its English original, the draft in your language, and a note on where it appears.',
    'If a draft reads well, leave its yellow cells empty. If it should change, write your wording under "Your correction"; add a reason under "Comment" if it helps.',
    'Keep everything in {curly braces} exactly as written: the game puts numbers and names there. You may move them within the sentence.',
    'Keep tags such as <kbd>Q</kbd> (key caps) and <b>Hint</b> (bold) as they are; text in <b>…</b> must match the button it names.',
    'Mathematical notation stays as it is: L*, s₀, s₁, …, C₀, and the Greek strategy names α to ψ.',
    'Rows marked "built from examples" are assembled by the program for different numbers. Correct the example sentences and the code will be adjusted to match.',
    'Please check the key terms below first; changing one of them usually changes many rows.',
    'To see the text in place, open the game and choose your language from the menu at the top of the page. Ask whoever sent you this file for the link.',
]


# the variety each draft targets, stated at the top of its workbook
VARIETY = {
    'es': 'This draft targets neutral Latin American Spanish, as used in Mexico, rather than the Spanish of Spain.',
    'pt': 'This draft targets Brazilian Portuguese.',
}


def write(code, lang):
    rows = {r['key']: r for r in lang['rows']}
    wb = Workbook()

    ws = wb.active
    ws.title = 'How to review'
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 44
    ws.column_dimensions['C'].width = 44
    ws['A1'] = f'Square Packing · review of the {lang["name"]} draft'
    ws['A1'].font = TITLE
    r = 3
    for line in ([VARIETY[code]] if code in VARIETY else []) + INSTRUCTIONS:
        ws.cell(row=r, column=1, value=line).font = BODY
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        ws.cell(row=r, column=1).alignment = WRAP
        ws.row_dimensions[r].height = 30
        r += 1

    r += 1
    for c, h in enumerate(['Key term', 'English', f'Draft ({lang["name"]})'], 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font, cell.fill = HEAD, HEAD_FILL
    r += 1
    for term, key in TERMS:
        if key is None:
            en = ', '.join(rows[k]['en'] for k in rows if k.startswith('strat.'))
            dr = ', '.join(rows[k]['draft'] for k in rows if k.startswith('strat.'))
        else:
            en, dr = rows[key]['en'], rows[key]['draft']
        for c, v in enumerate([term, en, dr], 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.alignment, cell.border = BODY, WRAP, THIN
        r += 1

    r += 1
    ws.cell(row=r, column=1, value='Example of a filled-in row (illustration only, not a real correction)').font = BOLD
    r += 1
    ex = rows['btn.shuffle']
    for c, (h, v) in enumerate([('English', ex['en']), ('Your correction', '[your wording, only if the draft should change]'),
                                ('Comment', '[optional: why, or which other rows it affects]')], 1):
        ws.cell(row=r, column=c, value=h).font = BOLD
        cell = ws.cell(row=r + 1, column=c, value=v)
        cell.font, cell.alignment = BODY, WRAP
        if c > 1:
            cell.fill = INPUT_FILL
    ws.cell(row=r + 3, column=1, value='Yellow cells are the only ones to edit.').font = MUTED

    st = wb.create_sheet('Strings')
    heads = ['#', 'Key', 'Where it appears', 'English', f'Draft ({lang["name"]})', 'Your correction', 'Comment']
    widths = [5, 18, 40, 46, 46, 40, 30]
    for c, (h, w) in enumerate(zip(heads, widths), 1):
        cell = st.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.alignment = HEAD, HEAD_FILL, WRAP
        st.column_dimensions[cell.column_letter].width = w
    st.freeze_panes = 'C2'
    for i, row in enumerate(lang['rows'], 2):
        context = row['context'] + ('  [built from examples]' if row['built'] else '')
        vals = [row['n'], row['key'], context, row['en'], row['draft'], None, None]
        for c, v in enumerate(vals, 1):
            cell = st.cell(row=i, column=c, value=v)
            cell.font = MUTED if c in (2, 3) else BODY
            cell.alignment = WRAP
            cell.border = THIN
            if c >= 6:
                cell.fill = INPUT_FILL
        lines = max(len(str(row['en'])) // 42 + str(row['en']).count('\n') + 1,
                    len(str(row['draft'])) // 42 + str(row['draft']).count('\n') + 1,
                    len(context) // 38 + 1)
        st.row_dimensions[i].height = max(15, 13 * lines)
    st.auto_filter.ref = f'A1:G{len(lang["rows"]) + 1}'

    out = ROOT / 'release' / 'review' / f'square-packing-review-{code}.xlsx'
    wb.save(out)
    return out


for code, lang in data.items():
    print(write(code, lang).name)
